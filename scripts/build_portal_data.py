import os
import re
import csv
import json
import time
import glob
import shutil
import hashlib
import sqlite3
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed
import docx
from docx.oxml.ns import qn
from PIL import Image

from render_helpers import (
    ROOT_DIR,
    PORTAL_DIR,
    STAGING_DIR,
    DATA_DIR,
    CACHE_PATH,
    CANONICAL_ROLES,
    sanitize_roles,
    extract_canonical_actors,
    render_screen_mockup,
    render_sequence_diagram_png,
    render_class_diagram_36_png,
)

# Load .env
ENV = {}
env_file = os.path.join(PORTAL_DIR, ".env")
if os.path.exists(env_file):
    for line in open(env_file, encoding="utf-8"):
        if "=" in line and not line.strip().startswith("#"):
            k, v = line.strip().split("=", 1)
            ENV[k.strip()] = v.strip()

CLOUD_NAME = ENV.get("Cloudinary__CloudName", "dvwt6npcl")
API_KEY = ENV.get("Cloudinary__ApiKey", "")
API_SECRET = ENV.get("Cloudinary__ApiSecret", "")

# Load Cloudinary upload cache
if os.path.exists(CACHE_PATH):
    try:
        UPLOAD_CACHE = json.load(open(CACHE_PATH, "r", encoding="utf-8"))
    except Exception:
        UPLOAD_CACHE = {}
else:
    UPLOAD_CACHE = {}

def upload_to_cloudinary(local_path: str, folder: str, public_id: str) -> dict:
    file_bytes = open(local_path, "rb").read()
    md5_hash = hashlib.md5(file_bytes).hexdigest()
    cache_key = f"{folder}/{public_id}"
    if cache_key in UPLOAD_CACHE and UPLOAD_CACHE[cache_key].get("md5") == md5_hash:
        return UPLOAD_CACHE[cache_key]

    ts = str(int(time.time()))
    overwrite = "true"
    sig_str = f"folder={folder}&overwrite={overwrite}&public_id={public_id}&timestamp={ts}{API_SECRET}"
    sig = hashlib.sha1(sig_str.encode("utf-8")).hexdigest()

    for attempt in range(3):
        try:
            with open(local_path, "rb") as f:
                r = requests.post(
                    f"https://api.cloudinary.com/v1_1/{CLOUD_NAME}/image/upload",
                    data={
                        "api_key": API_KEY,
                        "timestamp": ts,
                        "folder": folder,
                        "public_id": public_id,
                        "overwrite": overwrite,
                        "signature": sig,
                    },
                    files={"file": f},
                    timeout=60,
                )
            if r.status_code == 200:
                data = r.json()
                res = {
                    "url": data["secure_url"],
                    "width": data.get("width", 1920),
                    "height": data.get("height", 1080),
                    "format": data.get("format", "png"),
                    "bytes": data.get("bytes", len(file_bytes)),
                    "md5": md5_hash,
                }
                UPLOAD_CACHE[cache_key] = res
                return res
        except Exception as e:
            print(f"Retry {attempt+1} for {public_id}: {e}")
            time.sleep(1.5)
    raise RuntimeError(f"Failed to upload {local_path} to Cloudinary")

# 1. Load Vietnamese CSV descriptions
csv_path = os.path.join(ROOT_DIR, "Các FR của hệ thống  - Feature_Function Requirement.csv")
csv_fr_map = {}
if os.path.exists(csv_path):
    with open(csv_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 6:
                code = row[0].strip()
                if re.match(r"^3\.(?:[2-9]|10)\.\d+$", code):
                    csv_fr_map[code] = {
                        "viDescription": sanitize_roles(row[5].strip()),
                        "docAuthor": row[6].strip() if len(row) > 6 else "",
                        "codeAuthor": row[7].strip() if len(row) > 7 else "",
                        "note": row[8].strip() if len(row) > 8 else "",
                    }

# 2. Parse Report 3 Official DOCX
r3_path = os.path.join(ROOT_DIR, "brandhub-final_report", "report3", "report3_official.docx")
r3_doc = docx.Document(r3_path)

# Parse Tables from Report 3
# Table 1: Actors
actors_table = []
for row in r3_doc.tables[1].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 3:
        actors_table.append({
            "index": cells[0],
            "role": sanitize_roles(cells[1]),
            "description": sanitize_roles(cells[2]),
        })

# Table 2: 141 Use Cases
use_cases_table = []
for row in r3_doc.tables[2].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 4:
        uc_id = cells[0]
        uc_num = int(re.sub(r"\D", "", uc_id)) if re.sub(r"\D", "", uc_id) else 0
        use_cases_table.append({
            "id": uc_id,
            "name": cells[1],
            "actors": extract_canonical_actors(cells[2]),
            "actorsRaw": sanitize_roles(cells[2]),
            "description": sanitize_roles(cells[3]),
            "ucNum": uc_num,
        })

# Table 3: Screen Descriptions
screen_descriptions = []
for row in r3_doc.tables[3].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 4:
        screen_descriptions.append({
            "index": cells[0],
            "feature": cells[1],
            "screen": cells[2],
            "description": sanitize_roles(cells[3]),
        })

# Table 4: Screen Authorization
screen_auth = []
for row in r3_doc.tables[4].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 7:
        screen_auth.append({
            "screen": cells[0],
            "Admin": cells[1],
            "Owner": cells[2],
            "Manager": cells[3],
            "Creator": cells[4],
            "Client": cells[5],
            "Guest": cells[6],
        })

# Table 5: Non-Screen System Functions
non_screen_funcs = []
for row in r3_doc.tables[5].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 4:
        non_screen_funcs.append({
            "index": cells[0],
            "feature": cells[1],
            "systemFunction": cells[2],
            "description": sanitize_roles(cells[3]),
        })

# Table 6: Entities
entities_table = []
for row in r3_doc.tables[6].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 3:
        entities_table.append({
            "index": cells[0],
            "entity": cells[1],
            "description": sanitize_roles(cells[2]),
        })

# Table 7: 65 Business Rules
business_rules = []
br_lookup = {}
for row in r3_doc.tables[7].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 2 and cells[0]:
        item = {
            "code": cells[0],
            "definition": sanitize_roles(cells[1]),
            "relatedFrs": [],
        }
        business_rules.append(item)
        br_lookup[cells[0]] = item

# Table 8: 60 Application Messages
messages_list = []
msg_lookup = {}
for row in r3_doc.tables[8].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    if len(cells) >= 5 and cells[1]:
        item = {
            "index": cells[0],
            "code": cells[1],
            "type": cells[2],
            "context": sanitize_roles(cells[3]),
            "content": sanitize_roles(cells[4]),
            "relatedFrs": [],
        }
        messages_list.append(item)
        msg_lookup[cells[1]] = item

# Domain map
DOMAINS = [
    {"code": "3.2", "title": "Authentication & Authorization", "frCount": 9, "summary": "User registration, email/Google OAuth sign-in, password reset & change, 6-digit OTP verification, TOTP 2FA, session logout, and soft account deletion."},
    {"code": "3.3", "title": "User & Client Profile", "frCount": 4, "summary": "Personal user profile management, avatar S3 storage, and reusable cross-workspace Client brand profile management."},
    {"code": "3.4", "title": "Agency & Workspace", "frCount": 26, "summary": "Multi-tenant Agency lifecycle, member invitations (72h expiry), Workspace creation with mandatory Manager assignment, IANA timezone configuration, and Workspace templates."},
    {"code": "3.5", "title": "Media Package & Media Campaign", "frCount": 13, "summary": "Template & custom media packages, two-party package negotiation & approval, detailed media campaign planning, automated task backlog generation, and ad-hoc content requests."},
    {"code": "3.6", "title": "Content & Workflow", "frCount": 35, "summary": "Task backlog & detail specification, 4 task views (List, Calendar, Gantt, Kanban), multi-stage approval sequence (Creator -> Manager -> Client), Unicode content editor, Material Repository with Google Drive sync & watermarking, Hashtags, Surveys, Mail Templates, and AI Compliance check."},
    {"code": "3.7", "title": "AI Features", "frCount": 19, "summary": "Trending topics & Neo4j Knowledge Graph virality scoring, ChromaDB Brand RAG documents, AI caption/image/video/Virtual Ambassador generation, hashtag trend suggestions, and third-party collaborator recommendations."},
    {"code": "3.8", "title": "Publishing & Social", "frCount": 15, "summary": "OAuth 2.0 / PKCE social account linking, post analytics dashboard, post preview & scheduling, and automated publishing across 4 platforms: Facebook, Instagram, TikTok, and Threads."},
    {"code": "3.9", "title": "Subscription & Billing", "frCount": 6, "summary": "Subscription plan upgrade flow, SePay / PayOS (VietQR) automated webhook payment verification, invoice history, and workspace/creator AI Credit tracking & top-up."},
    {"code": "3.10", "title": "Admin Functions", "frCount": 14, "summary": "System-wide push notifications, platform KPI statistics, real-time server/service health monitoring, content moderation queue, user strike & status governance, AI Credit price catalog & CSV import, and revenue PDF reporting."},
]
domain_title_map = {d["code"]: d["title"] for d in DOMAINS}

# Fallback screenshot mapping for FRs that don't have an embedded blip inside report3_official.docx
SCREENSHOT_FALLBACKS = {
    "3.4.15": r"brandhub-infrastructure\docs\feature\agency-workspace\3-4-13-view-workspace-profile\screenshot.png",
    "3.4.20": r"brandhub-infrastructure\docs\feature\agency-workspace\3-4-17-save-workspace-template\screenshot.png",
    "3.5.6": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-19.png",
    "3.6.1": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-01.png",
    "3.6.2": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-02.png",
    "3.6.3": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-20.png",
    "3.6.4": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-03.png",
    "3.6.5": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-06.png",
    "3.6.6": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-07.png",
    "3.6.7": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-11.png",
    "3.6.8": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-10.png",
    "3.6.9": r"brandhub-infrastructure\docs\screenshots\D16\DA-D16-15.png",
    "3.6.10": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-01.png",
    "3.6.11": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-09.png",
    "3.6.34": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-06.png",
    "3.7.2": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-27.png",
    "3.7.3": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-24.png",
    "3.7.4": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-20.png",
    "3.7.8": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-08.png",
    "3.7.12": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-15.png",
    "3.7.14": r"brandhub-infrastructure\docs\screenshots\D17\DA-D17-28.png",
    "3.8.7": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-12.png",
    "3.8.8": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-06.png",
    "3.8.9": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-07.png",
    "3.8.10": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-08.png",
    "3.8.11": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-17.png",
    "3.8.12": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-18.png",
    "3.8.13": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-19.png",
    "3.8.14": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-13.png",
    "3.8.15": r"brandhub-infrastructure\docs\screenshots\D18\DA-D18-20.png",
    "3.9.1": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-01.png",
    "3.9.2": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-02.png",
    "3.9.3": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-03.png",
    "3.9.4": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-04.png",
    "3.9.5": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-05.png",
    "3.10.3": r"brandhub-infrastructure\docs\screenshots\D19\DA-D19-14.png",
}

# Parse all 141 FR sections in report3_official.docx
fr_heading_indices = []
for idx, p in enumerate(r3_doc.paragraphs):
    s = p.style.name if p.style else ""
    t = p.text.strip()
    m = re.match(r"^(3\.(?:[2-9]|10)\.\d+)\s+(.*)$", t)
    if s == "Heading 4" and m:
        fr_heading_indices.append((idx, m.group(1), m.group(2).strip()))

print(f"Found {len(fr_heading_indices)} FR sections in report3_official.docx")

all_figures = []
all_frs = []

# First, extract Report 3 Overview Figures (Figure 1 to Figure 9)
r3_overview_meta = [
    {
        "figNum": 1,
        "id": "r3-fig-1",
        "title": "System Context Diagram",
        "caption": "Figure 1 - System Context Diagram",
        "section": "1.2 System Context",
        "category": "Architecture & Context",
        "actors": ["Guest", "Owner", "Manager", "Creator", "Client", "Admin"],
        "summary": "Illustrates the boundary of the BrandHub platform, how the 6 canonical user actors interact with the system, and how BrandHub integrates with external AI providers (Groq LLM, Stability AI, Google Veo), social publishing networks (Facebook, Instagram, TikTok, Threads), Google Drive Storage, and SePay / PayOS payment webhooks.",
        "components": [
            {"name": "BrandHub Core Platform", "type": "System Boundary", "description": "Central multi-tenant SaaS platform managing Agencies, Workspaces, Media Campaigns, Tasks, AI Generation, and Social Publishing."},
            {"name": "6 Canonical Actors", "type": "Actors", "description": "Guest, Owner, Manager, Creator, Client, and Admin interacting via HTTPS / REST / WebSocket."},
            {"name": "AI Model Providers", "type": "External Service", "description": "Groq Cloud LLM, Stability AI / InstantID, and Google Veo for caption, image, ambassador, and video synthesis."},
            {"name": "4 Social Media APIs", "type": "External Service", "description": "Facebook Graph API, Instagram Graph API, TikTok Content Posting API (PKCE), and Threads Publishing API."},
            {"name": "SePay / PayOS Gateway", "type": "External Service", "description": "VietQR payment order generation and real-time webhook settlement verification."},
        ],
    },
    {
        "figNum": 2,
        "id": "r3-fig-2",
        "title": "System Use Case Diagram",
        "caption": "Figure 2 - System Use Case Diagram",
        "section": "2.2.1 Use Case Diagram(s)",
        "category": "Use Case Diagram",
        "actors": ["Guest", "Owner", "Manager", "Creator", "Client", "Admin"],
        "summary": "Comprehensive Use Case Overview mapping all 141 atomic CRUD Use Cases (UC-001 to UC-141) across the 9 functional domains and 6 canonical system actors, highlighting <<include>> and <<extend>> relationships.",
        "components": [
            {"name": "Authentication & Profile (UC-001 – UC-013)", "type": "Use Case Package", "description": "Account registration, email/Google OAuth sign-in, OTP/2FA verification, and User/Client profile management."},
            {"name": "Agency & Workspace (UC-014 – UC-039)", "type": "Use Case Package", "description": "Agency creation, 72-hour member invitations, Workspace creation with mandatory Manager, timezone & template governance."},
            {"name": "Media Package & Campaign (UC-040 – UC-052)", "type": "Use Case Package", "description": "Package negotiation, two-party approval, campaign planning, auto-generated task backlog, and content requests."},
            {"name": "Content, AI, Publishing & Billing (UC-053 – UC-141)", "type": "Use Case Package", "description": "Task workflow, 3-tier approval sequence, AI studio, 4-platform social publishing, SePay billing, and Admin governance."},
        ],
    },
    {
        "figNum": 3,
        "id": "r3-fig-3",
        "title": "Screen Flow for Guest",
        "caption": "Figure 3 - Screen Flow for Guest",
        "section": "3.1.1.1 Screen flow for Guest",
        "category": "Screen Flow",
        "actors": ["Guest"],
        "summary": "Navigation flow for unauthenticated Guest users across the Landing Page, Sign Up (/register), Sign In (/login), Google OAuth flow, OTP Verification (/verify-otp), and Password Reset (/forgot-password).",
        "components": [
            {"name": "Landing Page (/)", "type": "Entry Screen", "description": "Public product introduction, pricing catalog preview, and authentication entry points."},
            {"name": "Register & Login (/register, /login)", "type": "Auth Screens", "description": "Email/password registration and authentication or 1-click Google OAuth 2.0."},
            {"name": "Verify OTP & Reset Password", "type": "Verification Screens", "description": "6-digit single-use OTP verification (10-minute validity) for account activation and password recovery."},
        ],
    },
    {
        "figNum": 4,
        "id": "r3-fig-4",
        "title": "Screen Flow for Client",
        "caption": "Figure 4 - Screen Flow for Client",
        "section": "3.1.1.2 Screen flow for Client",
        "category": "Screen Flow",
        "actors": ["Client"],
        "summary": "Dedicated navigation flow for Client users within the Client Portal: managing reusable Client Brand Profiles, negotiating & approving Media Packages and Media Campaigns, submitting ad-hoc Content Requests, and performing final Client approval on Tasks.",
        "components": [
            {"name": "Client Profile (/client/profile)", "type": "Profile Screen", "description": "Reusable brand identity profile synchronized across all joined Workspaces."},
            {"name": "Package & Campaign Approval", "type": "Negotiation Screen", "description": "Two-party review, counter-proposal, and final approval of Media Packages and Campaigns."},
            {"name": "Task Review & Content Request", "type": "Collaboration Screen", "description": "Final approval/rejection of deliverables and submission of supplemental Content Requests."},
        ],
    },
    {
        "figNum": 5,
        "id": "r3-fig-5",
        "title": "Screen Flow for Admin",
        "caption": "Figure 5 - Screen Flow for Admin",
        "section": "3.1.1.3 Screen flow for Admin",
        "category": "Screen Flow",
        "actors": ["Admin"],
        "summary": "System-level administration flow for Admin users covering Platform Statistics Overview, Real-time Server & Service Health Monitoring, Content Moderation Queue, User Account & Violation Governance, AI Credit Price Catalog, and Revenue PDF Reporting.",
        "components": [
            {"name": "Platform Statistics & Health Monitor", "type": "Observability Console", "description": "Macro KPIs and real-time CPU/RAM/Disk/Database/RabbitMQ health telemetry."},
            {"name": "Content Moderation & User Governance", "type": "Compliance Console", "description": "Reviewing AI-flagged posts, issuing strikes, and managing account activation/sanctions."},
            {"name": "AI Credit Pricing & Revenue Dashboard", "type": "Billing Console", "description": "Configuring unit credit costs (manual & CSV import), viewing SePay receipts, and exporting PDF reports."},
        ],
    },
    {
        "figNum": 6,
        "id": "r3-fig-6",
        "title": "Screen Flow for Owner",
        "caption": "Figure 6 - Screen Flow for Owner",
        "section": "3.1.1.4 Screen flow for Owner",
        "category": "Screen Flow",
        "actors": ["Owner"],
        "summary": "End-to-end navigation flow for Agency Owners: creating and managing Agencies, inviting Agency members, provisioning Workspaces with assigned Managers, overseeing Campaigns and Tasks, and managing Subscription Plans, SePay Payments, and AI Credits.",
        "components": [
            {"name": "Agency Directory & Dashboard", "type": "Organization Layer", "description": "Creating agencies, tracking cross-workspace KPIs, and inviting agency members."},
            {"name": "Workspace & Template Governance", "type": "Workspace Layer", "description": "Creating workspaces, assigning mandatory Managers, and managing reusable workspace templates."},
            {"name": "Subscription, SePay Billing & AI Credits", "type": "Billing Layer", "description": "Upgrading subscription tiers, paying via VietQR (SePay/PayOS), and allocating AI credits."},
        ],
    },
    {
        "figNum": 7,
        "id": "r3-fig-7",
        "title": "Screen Flow for Manager",
        "caption": "Figure 7 - Screen Flow for Manager",
        "section": "3.1.1.5 Screen flow for Manager",
        "category": "Screen Flow",
        "actors": ["Manager"],
        "summary": "Operational workflow for Workspace Managers: configuring Workspace profiles & timezones, negotiating Media Packages & Campaigns with Clients, identifying Task details, assigning Tasks to Creators, reviewing content in the 3-tier approval chain, and scheduling/publishing posts across 4 social platforms.",
        "components": [
            {"name": "Workspace Operations & Members", "type": "Management Screen", "description": "Managing workspace members, roles, timezone configuration, and templates."},
            {"name": "Campaign to Task Backlog Pipeline", "type": "Planning Screen", "description": "Translating approved Campaigns and Content Requests into detailed, assigned Creator tasks."},
            {"name": "Approval Review & Multi-Platform Publishing", "type": "Execution Screen", "description": "Manager-level quality review and publishing/scheduling to Facebook, Instagram, TikTok, and Threads."},
        ],
    },
    {
        "figNum": 8,
        "id": "r3-fig-8",
        "title": "Screen Flow for Creator",
        "caption": "Figure 8 - Screen Flow for Creator",
        "section": "3.1.1.6 Screen flow for Creator",
        "category": "Screen Flow",
        "actors": ["Creator"],
        "summary": "Creative production flow for Creators: viewing assigned Tasks across List/Calendar/Gantt/Kanban boards, drafting posts in the Unicode Content Editor, using AI Studio (captions, images, videos, virtual ambassadors), managing the Material Repository & Hashtag Collections, and submitting work for review.",
        "components": [
            {"name": "Multi-View Task Board", "type": "Task Screen", "description": "Tracking assigned tasks via List, Calendar, Gantt Timeline, and Kanban views."},
            {"name": "Content Writing & AI Studio", "type": "Creation Screen", "description": "Rich Unicode text editor integrated with AI Caption, Image, Video, and Ambassador generation."},
            {"name": "Material Repository & Watermark Tool", "type": "Asset Screen", "description": "Uploading RAW/Retouched assets, syncing Google Drive, applying brand watermarks, and inserting hashtag sets."},
        ],
    },
    {
        "figNum": 9,
        "id": "r3-fig-9",
        "title": "Entity Relationship Diagram",
        "caption": "Figure 9 - Entity Relationship Diagram",
        "section": "3.1.5 Entity Relationship Diagram",
        "category": "Database & ERD",
        "actors": ["Guest", "Owner", "Manager", "Creator", "Client", "Admin"],
        "summary": "Complete polyglot persistence Entity Relationship Diagram covering 28 core entities across PostgreSQL (Users, Agencies, Workspaces, Subscriptions, Invoices, AI Credit Ledgers) and MongoDB / Neo4j / ChromaDB (Media Packages, Campaigns, Tasks, Approvals, Materials, Social Posts, Trend Knowledge Graph).",
        "components": [
            {"name": "Identity & Organization Entities", "type": "PostgreSQL", "description": "User, ClientProfile, Agency, AgencyMember, AgencyInvitation, Workspace, WorkspaceMember, WorkspaceTemplate."},
            {"name": "Campaign & Workflow Entities", "type": "Document / Relational", "description": "MediaPackage, WorkspaceMediaPackage, MediaCampaign, ContentRequest, Task, TaskApproval, TaskComment, ContentVersion."},
            {"name": "AI, Social & Billing Entities", "type": "Polyglot Storage", "description": "MaterialRepository, HashtagCollection, BrandDocument, CandidateTrend, SocialAccount, ScheduledPost, SubscriptionPlan, PaymentTransaction, AiCreditLedger."},
        ],
    },
]

# Extract Figure 1 - 9 images from report3_official.docx
for idx, p in enumerate(r3_doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r"^Figure\s+([1-9])\s*-\s*(.*)$", txt)
    if m:
        fnum = int(m.group(1))
        rids = []
        j = idx - 1
        while j >= 0 and j >= idx - 12:
            blips = r3_doc.paragraphs[j]._element.xpath(".//a:blip")
            if blips:
                for b in blips:
                    rids.append(b.get(qn("r:embed")))
            elif rids:
                break
            j -= 1
        rids = list(reversed(rids))
        meta = next(x for x in r3_overview_meta if x["figNum"] == fnum)
        local_img_path = os.path.join(STAGING_DIR, f"r3_fig_{fnum}.png")
        sub_images_local = []
        if len(rids) == 1:
            part = r3_doc.part.related_parts[rids[0]]
            open(local_img_path, "wb").write(part.blob)
        elif len(rids) > 1:
            # Stitch the 9 Use Case diagrams into a crisp 3x3 grid HD image AND keep sub-images
            pil_imgs = []
            for s_idx, rid in enumerate(rids):
                part = r3_doc.part.related_parts[rid]
                sub_p = os.path.join(STAGING_DIR, f"r3_fig_{fnum}_part_{s_idx+1}.png")
                open(sub_p, "wb").write(part.blob)
                sub_images_local.append((f"r3_fig_{fnum}_part_{s_idx+1}", DOMAINS[s_idx]["code"] + " " + DOMAINS[s_idx]["title"] if s_idx < len(DOMAINS) else f"Part {s_idx+1}", sub_p))
                pil_imgs.append(Image.open(sub_p).convert("RGB"))
            # Stitch 3x3 grid
            cols = 3
            rows = (len(pil_imgs) + cols - 1) // cols
            cell_w = max(im.width for im in pil_imgs)
            cell_h = max(im.height for im in pil_imgs)
            canvas = Image.new("RGB", (cols * cell_w + (cols + 1) * 24, rows * cell_h + (rows + 1) * 24), (255, 255, 255))
            for idx_im, im in enumerate(pil_imgs):
                r_i = idx_im // cols
                c_i = idx_im % cols
                ox = 24 + c_i * (cell_w + 24) + (cell_w - im.width) // 2
                oy = 24 + r_i * (cell_h + 24) + (cell_h - im.height) // 2
                canvas.paste(im, (ox, oy))
            canvas.save(local_img_path, "PNG", optimize=True)

        meta["localPath"] = local_img_path
        meta["subImagesLocal"] = sub_images_local
        meta["reportId"] = "report3"
        meta["reportName"] = "Report 3 — Software Requirement Specification (SRS)"
        all_figures.append(meta)

# Now parse all 141 FRs (Figure 10 to Figure 150)
for i, (start_idx, fr_code, fr_title) in enumerate(fr_heading_indices):
    end_idx = fr_heading_indices[i + 1][0] if i + 1 < len(fr_heading_indices) else 4806
    dom_code = ".".join(fr_code.split(".")[:2])
    dom_title = domain_title_map.get(dom_code, "")
    uc_obj = use_cases_table[i] if i < len(use_cases_table) else {"id": f"UC-{i+1:03d}", "name": fr_title, "actors": ["Owner"], "description": ""}
    uc_obj["frCode"] = fr_code

    trigger = ""
    actors_raw = ""
    purpose = ""
    interface_desc = ""
    data_processing = ""
    fig_num = 10 + i
    fig_caption = f"Figure {fig_num} - {fr_title}"
    rids = []
    data_specs = {"inputRequired": "None.", "inputOptional": "None.", "systemData": "None.", "output": ""}
    fr_brs = []
    validations = []
    normal_flow = []
    abnormal_cases = []
    post_conditions = []

    mode = None
    for j in range(start_idx + 1, end_idx):
        p = r3_doc.paragraphs[j]
        blips = p._element.xpath(".//a:blip")
        for b in blips:
            rids.append(b.get(qn("r:embed")))
        txt = sanitize_roles(p.text.strip())
        if not txt:
            continue
        if txt.startswith("3.") and p.style and p.style.name == "Heading 3":
            break

        if txt == "Function Trigger":
            mode = "trigger"
            continue
        elif txt == "Function Description":
            mode = "desc"
            continue
        elif txt == "Screen Layout":
            mode = "layout"
            continue
        elif txt == "Function Details":
            mode = "details"
            continue
        elif txt == "Data Specifications":
            mode = "data"
            continue
        elif txt == "Business Rules":
            mode = "br"
            continue
        elif txt == "Validation":
            mode = "val"
            continue
        elif txt == "Functionalities":
            mode = "func"
            continue
        elif txt == "Normal Flow":
            mode = "normal"
            continue
        elif txt in ("Abnormal Cases", "Exception Flow"):
            mode = "abnormal"
            continue
        elif txt == "Post-Conditions":
            mode = "post"
            continue

        if re.match(r"^Figure\s+\d+\s*-", txt):
            fig_caption = txt
            m_fn = re.match(r"^Figure\s+(\d+)", txt)
            if m_fn:
                fig_num = int(m_fn.group(1))
            continue

        if mode == "trigger":
            trigger = (trigger + " " + txt).strip()
        elif mode == "desc":
            if txt.startswith("Actors / Roles:"):
                actors_raw = txt.replace("Actors / Roles:", "").strip().rstrip(".")
            elif txt.startswith("Purpose:"):
                purpose = txt.replace("Purpose:", "").strip()
            elif txt.startswith("Interface:"):
                interface_desc = txt.replace("Interface:", "").strip()
            elif txt.startswith("Data Processing:"):
                data_processing = txt.replace("Data Processing:", "").strip()
        elif mode == "data":
            if txt.startswith("Input required:"):
                data_specs["inputRequired"] = txt.replace("Input required:", "").strip()
            elif txt.startswith("Input optional:"):
                data_specs["inputOptional"] = txt.replace("Input optional:", "").strip()
            elif txt.startswith("System data:"):
                data_specs["systemData"] = txt.replace("System data:", "").strip()
            elif txt.startswith("Output:"):
                data_specs["output"] = txt.replace("Output:", "").strip()
        elif mode == "br":
            m_br = re.match(r"^(BR-\d+)\s*:\s*(.*)$", txt)
            if m_br:
                fr_brs.append({"code": m_br.group(1), "text": m_br.group(2)})
            elif txt.lower() != "none.":
                fr_brs.append({"code": "BR", "text": txt})
        elif mode == "val":
            if txt.lower() != "none.":
                validations.append(txt)
        elif mode == "normal":
            normal_flow.append(re.sub(r"^\d+\.\s*", "", txt))
        elif mode == "abnormal":
            abnormal_cases.append(txt)
        elif mode == "post":
            post_conditions.append(txt)

    canonical_actors = extract_canonical_actors(actors_raw or uc_obj.get("actorsRaw", "Owner"))

    # Link BRs and MSGs across the whole FR text
    full_fr_text = " ".join([trigger, purpose, interface_desc, data_processing, json.dumps(data_specs), json.dumps(fr_brs), " ".join(validations), " ".join(normal_flow), " ".join(abnormal_cases)])
    linked_br_codes = sorted(set(re.findall(r"\bBR-\d{2}\b", full_fr_text)))
    linked_msg_codes = sorted(set(re.findall(r"\bMSG\d{2}\b", full_fr_text)))

    for b_code in linked_br_codes:
        if b_code in br_lookup and fr_code not in br_lookup[b_code]["relatedFrs"]:
            br_lookup[b_code]["relatedFrs"].append(fr_code)
    for m_code in linked_msg_codes:
        if m_code in msg_lookup and fr_code not in msg_lookup[m_code]["relatedFrs"]:
            msg_lookup[m_code]["relatedFrs"].append(fr_code)

    # Determine local image for Figure (10..150)
    fig_id = f"r3-fig-{fig_num}"
    local_img_path = os.path.join(STAGING_DIR, f"r3_fig_{fig_num}_{fr_code.replace('.', '_')}.png")
    image_source_type = "Embedded SRS Screenshot"

    if rids:
        part = r3_doc.part.related_parts[rids[0]]
        open(local_img_path, "wb").write(part.blob)
    elif fr_code in SCREENSHOT_FALLBACKS and os.path.exists(os.path.join(ROOT_DIR, SCREENSHOT_FALLBACKS[fr_code])):
        shutil.copyfile(os.path.join(ROOT_DIR, SCREENSHOT_FALLBACKS[fr_code]), local_img_path)
        image_source_type = "Captured Application UI Screenshot"
    else:
        render_screen_mockup(
            local_img_path,
            fig_num,
            fr_code,
            fr_title,
            dom_title,
            canonical_actors,
            purpose,
            interface_desc,
            data_specs,
            fr_brs,
        )
        image_source_type = "High-Resolution UI Layout Blueprint"

    csv_info = csv_fr_map.get(fr_code, {})

    fig_entry = {
        "figNum": fig_num,
        "id": fig_id,
        "title": fr_title,
        "caption": fig_caption,
        "section": f"{fr_code} {fr_title}",
        "domainCode": dom_code,
        "domainTitle": dom_title,
        "frCode": fr_code,
        "ucId": uc_obj["id"],
        "category": "Screen Layout (SRS)",
        "imageSourceType": image_source_type,
        "actors": canonical_actors,
        "summary": purpose,
        "viSummary": csv_info.get("viDescription", ""),
        "interface": interface_desc,
        "dataProcessing": data_processing,
        "trigger": trigger,
        "dataSpecs": data_specs,
        "businessRules": fr_brs,
        "linkedBrCodes": linked_br_codes,
        "linkedMsgCodes": linked_msg_codes,
        "validations": validations,
        "normalFlow": normal_flow,
        "abnormalCases": abnormal_cases,
        "postConditions": post_conditions,
        "sequenceFigureId": f"seq-{fr_code}",
        "classDiagramFigureId": f"cd-{dom_code}",
        "localPath": local_img_path,
        "reportId": "report3",
        "reportName": "Report 3 — Software Requirement Specification (SRS)",
        "components": [
            {"name": "Target Route & Screen Shell", "type": "UI Boundary", "description": interface_desc or f"Screen interface for {fr_title}"},
            {"name": "Required & Optional Inputs", "type": "Form / Controls", "description": f"Required: {data_specs['inputRequired']} | Optional: {data_specs['inputOptional']}"},
            {"name": "Backend Data Processing", "type": "Service Logic", "description": data_processing},
            {"name": "Expected Output & Feedback", "type": "System Output", "description": data_specs["output"]},
        ],
    }
    all_figures.append(fig_entry)

    fr_entry = {
        "code": fr_code,
        "title": fr_title,
        "domainCode": dom_code,
        "domainTitle": dom_title,
        "ucId": uc_obj["id"],
        "ucName": uc_obj["name"],
        "ucDescription": uc_obj["description"],
        "figureId": fig_id,
        "figureNumber": fig_num,
        "figureCaption": fig_caption,
        "sequenceFigureId": f"seq-{fr_code}",
        "classDiagramFigureId": f"cd-{dom_code}",
        "actors": canonical_actors,
        "trigger": trigger,
        "purpose": purpose,
        "viSummary": csv_info.get("viDescription", ""),
        "docAuthor": csv_info.get("docAuthor", ""),
        "codeAuthor": csv_info.get("codeAuthor", ""),
        "interface": interface_desc,
        "dataProcessing": data_processing,
        "dataSpecs": data_specs,
        "businessRules": fr_brs,
        "linkedBrCodes": linked_br_codes,
        "linkedMsgCodes": linked_msg_codes,
        "validations": validations,
        "normalFlow": normal_flow,
        "abnormalCases": abnormal_cases,
        "postConditions": post_conditions,
    }
    all_frs.append(fr_entry)

print(f"Prepared {len(all_frs)} FRs and {len(all_figures)} Report 3 Figures.")

# 3. Extract Report 1 Figures
r1_path = os.path.join(ROOT_DIR, "brandhub-final_report", "report1", "Report1_Project Introduction.docx")
r1_doc = docx.Document(r1_path)
r1_figs_spec = [
    ("r1-fig-1", "rId8", "Figure 3.1 - Hootsuite Dashboard Interface", "3.1 Existing System — Hootsuite", "Competitor Analysis", "Illustrates Hootsuite's multi-column social media streams and scheduling interface. Analyzed in Report 1 to highlight its lack of structured Agency-to-Client Media Package negotiation and multi-tier approval workflows."),
    ("r1-fig-2", "rId10", "Figure 3.2 - Buffer Interface", "3.2 Existing System — Buffer", "Competitor Analysis", "Shows Buffer's queue-based social publishing calendar. While simple for solo creators, it lacks multi-tenant Agency/Workspace hierarchy, AI Brand RAG knowledge bases, and 3-tier Client approval chains."),
    ("r1-fig-3", "rId12", "Figure 3.3 - Jira Interface", "3.3 Existing System — Atlassian Jira", "Competitor Analysis", "Displays Jira's Kanban and issue tracking board. BrandHub adopts Jira's structured backlog and filtering strengths while adding native AI content generation and direct 4-platform social publishing."),
    ("r1-fig-4", "rId14", "Figure 3.4 - Notion Interface", "3.4 Existing System — Notion", "Competitor Analysis", "Shows Notion's document workspace used by agencies for manual campaign briefs. BrandHub replaces fragmented Notion docs with executable Media Campaigns that auto-generate trackable Tasks."),
    ("r1-fig-5", "rId15", "Figure 6.1 - BrandHub Multi-Channel Ecosystem Scope", "6.1 Project Scope & Major Features", "Project Scope", "High-level functional scope diagram showing BrandHub's 9 core modules connecting Agencies, Workspaces, AI Studio, and 4 social publishing platforms (Facebook, Instagram, TikTok, Threads)."),
]
for idx_r1, (fid, rid, cap, sec, cat, summ) in enumerate(r1_figs_spec, 1):
    if rid in r1_doc.part.related_parts:
        lp = os.path.join(STAGING_DIR, f"{fid.replace('-', '_')}.png")
        open(lp, "wb").write(r1_doc.part.related_parts[rid].blob)
        all_figures.append({
            "figNum": idx_r1,
            "id": fid,
            "title": cap.split(" - ", 1)[-1],
            "caption": cap,
            "section": sec,
            "category": cat,
            "actors": ["Guest", "Owner", "Manager", "Creator", "Client", "Admin"],
            "summary": summ,
            "localPath": lp,
            "reportId": "report1",
            "reportName": "Report 1 — Project Introduction",
            "components": [
                {"name": cap.split(" - ", 1)[-1], "type": cat, "description": summ},
            ],
        })

# 4. Extract Report 2 Figures
r2_path = os.path.join(ROOT_DIR, "brandhub-final_report", "report2", "Report2_Project Management Plan.docx")
r2_doc = docx.Document(r2_path)
r2_figs_spec = [
    ("r2-fig-1", "rId7", "Figure 1.1 - Work Breakdown Structure (WBS) & Estimation", "1.1 Scope & Estimation", "Project Management", "Hierarchical Work Breakdown Structure (WBS) across project phases, sprints, and the 9 functional domains covering all 141 atomic Functional Requirements."),
    ("r2-fig-2", "rId8", "Figure 2.1 - Iterative Agile-Scrum Lifecycle", "2.1 Project Process", "Project Management", "Illustrates the iterative Agile-Scrum software development lifecycle adopted by the 5-member BrandHub engineering team."),
    ("r2-fig-3", "rId9", "Figure 2.2 - Sprint Execution & AI Parallel Track", "2.1 Project Process", "Project Management", "Shows how 2-week Scrum Sprints synchronize between the Core Platform track (Spring Boot / React) and the AI Engineering track (FastAPI / RAG / Neo4j)."),
    ("r2-fig-4", "rId10", "Figure 2.3 - Scrum Master & Cross-Review Workflow", "2.1 Project Process", "Project Management", "Depicts the daily standup, sprint planning, peer code review, and mentor milestone review workflow."),
    ("r2-fig-5", "rId11", "Figure 2.4 - Multi-Stage Quality Gate & Testing Pipeline", "2.2 Quality Management", "Quality Assurance", "Shows the 5 testing levels: Static Analysis / Code Review, Unit Testing (2,287 cases), Integration Testing, System Testing (2,287 cases), and User Acceptance Testing (UAT)."),
    ("r2-fig-6", "rId19", "Figure 6.1 - Repository & Documentation Structure Diagram", "6.1 Document Management", "Configuration Management", "Shows the polyrepo GitHub organization structure (api-gateway, business-service, ai-service, publisher-service, web, mobile, infrastructure) and capstone documentation tree."),
    ("r2-fig-7", "rId20", "Figure 6.2 - Git Polyrepo Branching Strategy Diagram", "6.2 Source Code Management", "Configuration Management", "Illustrates the Git branching model (main, develop, feature/*, release/*, hotfix/*) with mandatory pull request reviews and CI verification."),
]
for idx_r2, (fid, rid, cap, sec, cat, summ) in enumerate(r2_figs_spec, 1):
    if rid in r2_doc.part.related_parts:
        lp = os.path.join(STAGING_DIR, f"{fid.replace('-', '_')}.png")
        open(lp, "wb").write(r2_doc.part.related_parts[rid].blob)
        all_figures.append({
            "figNum": idx_r2,
            "id": fid,
            "title": cap.split(" - ", 1)[-1],
            "caption": cap,
            "section": sec,
            "category": cat,
            "actors": ["Owner", "Manager", "Creator", "Admin"],
            "summary": summ,
            "localPath": lp,
            "reportId": "report2",
            "reportName": "Report 2 — Project Management Plan",
            "components": [
                {"name": cap.split(" - ", 1)[-1], "type": cat, "description": summ},
            ],
        })

# 5. Prepare 9 Domain Class Diagrams (cd-3.2 to cd-3.10)
cd_36_path = os.path.join(STAGING_DIR, "cd_3_6.png")
render_class_diagram_36_png(cd_36_path)

CD_SOURCES = {
    "3.2": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.2.png"),
    "3.3": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\sample_cd_33_streamlined.png"),
    "3.4": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\sample_cd_34_current.png"),
    "3.5": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.5.png"),
    "3.6": cd_36_path,
    "3.7": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.7.png"),
    "3.8": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.8.png"),
    "3.9": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.9.png"),
    "3.10": os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\all_cd_samples\3.10.png"),
}

# Extract class lists from BrandHub-Master-All-In-One.vpp
vpp_path = os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\master\BrandHub-Master-All-In-One.vpp")
vpp_cd_classes = {}
if os.path.exists(vpp_path):
    conn = sqlite3.connect(vpp_path)
    cur = conn.cursor()
    cur.execute("SELECT ID, NAME FROM DIAGRAM WHERE DIAGRAM_TYPE='ClassDiagram'")
    for d_id, d_name in cur.fetchall():
        m_dom = re.match(r"^(3\.\d+)", d_name.strip())
        if m_dom:
            d_code = m_dom.group(1)
            cur.execute("""
                SELECT DISTINCT me.NAME
                FROM DIAGRAM_ELEMENT de
                JOIN MODEL_ELEMENT me ON de.MODEL_ELEMENT_ID = me.ID
                WHERE de.DIAGRAM_ID = ? AND me.MODEL_TYPE = 'Class' AND me.NAME IS NOT NULL
                ORDER BY me.NAME
            """, (d_id,))
            vpp_cd_classes[d_code] = [r[0] for r in cur.fetchall() if r[0]]
    conn.close()

vpp_cd_classes["3.6"] = [
    "TaskController", "TaskService", "Task", "TaskApproval", "TaskComment", "ContentVersion",
    "MaterialRepositoryController", "MaterialService", "MaterialAsset", "HashtagCollection",
    "SurveyService", "MailTemplateService"
]

for idx_cd, dom in enumerate(DOMAINS, 1):
    d_code = dom["code"]
    d_title = dom["title"]
    fid = f"cd-{d_code}"
    src_p = CD_SOURCES[d_code]
    lp = os.path.join(STAGING_DIR, f"cd_{d_code.replace('.', '_')}.png")
    if src_p != lp:
        shutil.copyfile(src_p, lp)
    cls_list = vpp_cd_classes.get(d_code, [])
    comps = []
    for cname in cls_list[:18]:
        ctype = "Entity / Domain Model"
        if cname.endswith("Controller"):
            ctype = "<<Controller>> REST Boundary"
        elif cname.endswith("Service") or cname.endswith("ServiceImpl"):
            ctype = "<<Service>> Business Logic"
        elif cname.endswith("Repository"):
            ctype = "<<Repository>> Data Access"
        elif "Dto" in cname or "Request" in cname or "Response" in cname:
            ctype = "<<DTO>> Data Transfer"
        comps.append({
            "name": cname,
            "type": ctype,
            "description": f"Core {ctype.split(' ')[-1]} class in module {d_code} ({d_title}).",
        })
    all_figures.append({
        "figNum": idx_cd,
        "id": fid,
        "title": f"{d_code} - {d_title} Class Diagram",
        "caption": f"Class Diagram {d_code} - {d_title}",
        "section": f"Report 4 SDD — Section {d_code} Class Design",
        "domainCode": d_code,
        "domainTitle": d_title,
        "category": "UML Class Diagram",
        "actors": ["Owner", "Manager", "Creator", "Client", "Admin"],
        "summary": f"Detailed UML Domain & Service Class Diagram for module {d_code} ({d_title}), illustrating Controllers, Services, Repositories, Domain Entities, attributes, method signatures, and associations across {dom['frCount']} Functional Requirements.",
        "localPath": lp,
        "reportId": "report4",
        "reportName": "Report 4 — Software Design Document (SDD)",
        "components": comps,
    })

# 6. Map and Prepare all 141 Sequence Diagrams (seq-3.2.1 to seq-3.10.14)
CONSOLIDATED_DIR = os.path.join(ROOT_DIR, r"visual-paradigm-test\batch\inspect_consolidated")
PUML_DIR = os.path.join(ROOT_DIR, r"brandhub-final_report\visual-paradigm\test\transformed_pumls")

# Map canonical FR code -> old VP file prefix in inspect_consolidated & transformed_pumls
SEQ_CODE_MAP = {
    # 3.2 (1-to-1)
    "3.2.1": "3.2.1", "3.2.2": "3.2.2", "3.2.3": "3.2.3", "3.2.4": "3.2.4", "3.2.5": "3.2.5",
    "3.2.6": "3.2.6", "3.2.7": "3.2.7", "3.2.8": "3.2.8", "3.2.9": "3.2.9",
    # 3.3 (1-to-1)
    "3.3.1": "3.3.1", "3.3.2": "3.3.2", "3.3.3": "3.3.3", "3.3.4": "3.3.4",
    # 3.4 (26 FRs)
    "3.4.1": "3.4.1", "3.4.2": "3.4.2", "3.4.3": "3.4.3", "3.4.4": "3.4.4", "3.4.5": "3.4.5",
    "3.4.6": "3.4.6", "3.4.7": "3.4.7a", "3.4.8": "3.4.8a", "3.4.9": "3.4.7b", "3.4.10": "3.4.8b",
    "3.4.11": "3.4.9", "3.4.12": "3.4.10", "3.4.13": "3.4.11", "3.4.14": "3.4.12", "3.4.15": "3.4.13",
    "3.4.16": "3.4.14", "3.4.17": "3.4.15", "3.4.18": "3.4.16", "3.4.19": "3.4.17a", "3.4.20": "3.4.17b",
    "3.4.23": "3.4.18", "3.4.24": "3.4.19b", "3.4.25": "3.4.20", "3.4.26": "3.4.21",
    # 3.5 (13 FRs)
    "3.5.1": "3.5.1", "3.5.2": "3.5.2", "3.5.3": "3.5.3", "3.5.4": "3.5.4", "3.5.5": "3.5.5",
    "3.5.8": "3.5.6", "3.5.10": "3.5.7", "3.5.11": "3.5.8", "3.5.12": "3.5.9", "3.5.13": "3.5.10",
    # 3.6 (35 FRs)
    "3.6.1": "3.6.1", "3.6.2": "3.6.2", "3.6.3": "3.6.3", "3.6.4": "3.6.4", "3.6.5": "3.6.5",
    "3.6.6": "3.6.6", "3.6.7": "3.6.7", "3.6.8": "3.6.8", "3.6.9": "3.6.9",
    "3.6.11": "3.6.11", "3.6.12": "3.6.12", "3.6.13": "3.6.13", "3.6.14": "3.6.14",
    "3.6.19": "3.6.16", "3.6.20": "3.6.17", "3.6.21": "3.6.18", "3.6.22": "3.6.19",
    "3.6.23": "3.6.20", "3.6.24": "3.6.21", "3.6.25": "3.6.25", "3.6.26": "3.6.26",
    "3.6.27": "3.6.27", "3.6.28": "3.6.28", "3.6.29": "3.6.29", "3.6.30": "3.6.30",
    "3.6.31": "3.6.31", "3.6.32": "3.6.32", "3.6.33": "3.6.33", "3.6.34": "3.6.35",
    "3.6.35": "3.6.36",
    # 3.7 (19 FRs)
    "3.7.1": "3.7.1", "3.7.6": "3.7.2", "3.7.7": "3.7.3", "3.7.8": "3.7.4", "3.7.9": "3.7.5",
    "3.7.10": "3.7.6", "3.7.11": "3.7.7", "3.7.12": "3.7.8", "3.7.13": "3.7.9", "3.7.14": "3.7.10",
    "3.7.19": "3.7.12",
    # 3.8 (15 FRs)
    "3.8.1": "3.8.1a", "3.8.2": "3.8.2", "3.8.3": "3.8.3", "3.8.4": "3.8.4", "3.8.5": "3.8.5",
    "3.8.6": "3.8.6", "3.8.7": "3.8.7", "3.8.8": "3.8.8", "3.8.9": "3.8.9", "3.8.10": "3.8.10",
    "3.8.11": "3.8.11", "3.8.12": "3.8.12", "3.8.13": "3.8.13", "3.8.14": "3.8.14", "3.8.15": "3.8.15",
    # 3.9 (6 FRs — 3.9.2 Downgrade Plan is excluded per AGENTS.md!)
    "3.9.1": "3.9.1", "3.9.2": "3.9.3", "3.9.3": "3.9.4", "3.9.4": "3.9.5", "3.9.5": "3.9.6", "3.9.6": "3.9.7",
    # 3.10 (14 FRs)
    "3.10.1": "3.10.1", "3.10.2": "3.10.2", "3.10.3": "3.10.3a", "3.10.4": "3.10.4", "3.10.5": "3.10.5",
    "3.10.6": "3.10.6", "3.10.7": "3.10.7", "3.10.8": "3.10.8", "3.10.9": "3.10.9",
    "3.10.13": "3.10.11", "3.10.14": "3.10.12",
}

all_consolidated_pngs = glob.glob(os.path.join(CONSOLIDATED_DIR, "*.png"))
all_pumls = glob.glob(os.path.join(PUML_DIR, "*.puml"))

def find_vp_png(prefix: str):
    for p in all_consolidated_pngs:
        b = os.path.basename(p)
        if re.search(rf"___({re.escape(prefix)})_-_", b) or b.startswith(f"{prefix}_-_"):
            return p
    return None

def find_puml(prefix: str):
    for p in all_pumls:
        b = os.path.basename(p)
        if b.startswith(f"{prefix}_-_") or b.startswith(f"{prefix}.1_-_"):
            return p
    return None

def parse_puml_details(puml_path: str, fallback_actors: list, fr_title: str, normal_flow: list):
    lifelines = []
    messages = []
    alias_map = {}
    if puml_path and os.path.exists(puml_path):
        txt = open(puml_path, "r", encoding="utf-8", errors="ignore").read()
        for line in txt.splitlines():
            line_s = line.strip()
            m_part = re.match(r'^(actor|boundary|control|entity|database|collections|participant|queue)\s+"([^"]+)"\s+as\s+(\S+)', line_s)
            if m_part:
                kind, name, alias = m_part.groups()
                name = sanitize_roles(name)
                stereo_map = {
                    "actor": "Actor",
                    "boundary": "Boundary" if "UI" in name or "Screen" in name or "Page" in name or "Panel" in name or "Modal" in name or "View" in name else "Controller",
                    "control": "Service",
                    "entity": "Repository" if "Repository" in name or "Dao" in name else "Entity",
                    "database": "Database",
                    "collections": "External Service",
                    "queue": "Message Broker",
                }
                stereo = stereo_map.get(kind, "Component")
                alias_map[alias] = (len(lifelines), name)
                lifelines.append({
                    "name": name,
                    "stereotype": stereo,
                    "description": f"{stereo} participant coordinating {fr_title}.",
                })
                continue
            m_msg = re.match(r'^(\S+?)\s*(->|-->)\s*(\S+?)\s*:\s*(.+)$', line_s)
            if m_msg:
                src_a, arrow, dst_a, lbl = m_msg.groups()
                lbl = sanitize_roles(lbl.strip())
                src_idx, src_name = alias_map.get(src_a, (0, src_a))
                dst_idx, dst_name = alias_map.get(dst_a, (min(1, max(len(lifelines)-1, 0)), dst_a))
                messages.append({
                    "step": len(messages) + 1,
                    "from": src_name,
                    "to": dst_name,
                    "fromIdx": src_idx,
                    "toIdx": dst_idx,
                    "isReturn": arrow == "-->",
                    "label": lbl,
                })
    if not lifelines:
        act_name = fallback_actors[0] if fallback_actors else "Owner"
        clean_slug = re.sub(r"[^a-zA-Z0-9]", "", fr_title.title())[:18]
        lifelines = [
            {"name": act_name, "stereotype": "Actor", "description": f"Canonical {act_name} role initiating {fr_title}."},
            {"name": f"{clean_slug}Screen (UI)", "stereotype": "Boundary", "description": f"Frontend React boundary view for {fr_title}."},
            {"name": f"{clean_slug}Controller", "stereotype": "Controller", "description": f"REST API Controller handling endpoint requests for {fr_title}."},
            {"name": f"{clean_slug}Service", "stereotype": "Service", "description": f"Core domain service enforcing business rules & RBAC for {fr_title}."},
            {"name": f"{clean_slug}Repository", "stereotype": "Repository", "description": f"Persistence repository managing entity state for {fr_title}."},
            {"name": "Database", "stereotype": "Database", "description": "PostgreSQL / MongoDB persistence store."},
        ]
        for idx_step, nf in enumerate(normal_flow[:6]):
            messages.append({
                "step": len(messages) + 1,
                "from": lifelines[min(idx_step, 4)]["name"],
                "to": lifelines[min(idx_step + 1, 5)]["name"],
                "fromIdx": min(idx_step, 4),
                "toIdx": min(idx_step + 1, 5),
                "isReturn": False,
                "label": nf[:85],
            })
        messages.append({
            "step": len(messages) + 1,
            "from": lifelines[3]["name"],
            "to": lifelines[1]["name"],
            "fromIdx": 3,
            "toIdx": 1,
            "isReturn": True,
            "label": "200 OK (ApiResponse<DTO>) & update UI state",
        })
    return lifelines, messages

for idx_seq, fr in enumerate(all_frs, 1):
    fr_code = fr["code"]
    fr_title = fr["title"]
    fid = f"seq-{fr_code}"
    vp_prefix = SEQ_CODE_MAP.get(fr_code)
    vp_png = find_vp_png(vp_prefix) if vp_prefix else None
    puml_file = find_puml(vp_prefix) if vp_prefix else None

    lifelines, messages = parse_puml_details(puml_file, fr["actors"], fr_title, fr["normalFlow"])
    lp = os.path.join(STAGING_DIR, f"seq_{fr_code.replace('.', '_')}.png")

    if vp_png and os.path.exists(vp_png):
        shutil.copyfile(vp_png, lp)
    else:
        render_sequence_diagram_png(lp, fr_code, fr_title, fr["actors"], lifelines, messages)

    comps = [
        {"name": f"{lf['name']}", "type": f"<<{lf['stereotype']}>>", "description": lf["description"]}
        for lf in lifelines
    ]

    all_figures.append({
        "figNum": idx_seq,
        "id": fid,
        "title": f"Sequence Diagram — {fr_code} {fr_title}",
        "caption": f"Sequence Diagram {fr_code} - {fr_title}",
        "section": f"Report 4 SDD — {fr_code} {fr_title}",
        "domainCode": fr["domainCode"],
        "domainTitle": fr["domainTitle"],
        "frCode": fr_code,
        "ucId": fr["ucId"],
        "category": "UML Sequence Diagram",
        "actors": fr["actors"],
        "summary": f"End-to-end UML Sequence Diagram for FR {fr_code} ({fr_title}), detailing the synchronous/asynchronous interactions across {len(lifelines)} lifelines ({', '.join(l['name'] for l in lifelines[:5])}) and {len(messages)} message steps.",
        "viSummary": fr.get("viSummary", ""),
        "screenFigureId": fr["figureId"],
        "classDiagramFigureId": fr["classDiagramFigureId"],
        "lifelines": lifelines,
        "sequenceMessages": messages[:30],
        "normalFlow": fr["normalFlow"],
        "abnormalCases": fr["abnormalCases"],
        "linkedBrCodes": fr["linkedBrCodes"],
        "linkedMsgCodes": fr["linkedMsgCodes"],
        "localPath": lp,
        "reportId": "report4",
        "reportName": "Report 4 — Software Design Document (SDD)",
        "components": comps,
    })

print(f"Total figures staged for Cloudinary upload: {len(all_figures)}")

# 7. Upload all staged images (plus sub-images of r3-fig-2) to Cloudinary in parallel
upload_tasks = []
for fig in all_figures:
    folder = f"brandhub-docs-portal/{fig['reportId']}"
    pub_id = fig["id"].replace(".", "_")
    upload_tasks.append((fig["id"], fig["localPath"], folder, pub_id))
    for sub_id, sub_label, sub_lp in fig.get("subImagesLocal", []):
        upload_tasks.append((sub_id, sub_lp, folder, sub_id))

print(f"Uploading {len(upload_tasks)} image assets to Cloudinary ({CLOUD_NAME})...")
results_map = {}

with ThreadPoolExecutor(max_workers=10) as executor:
    future_to_key = {
        executor.submit(upload_to_cloudinary, lp, folder, pub_id): key
        for key, lp, folder, pub_id in upload_tasks
    }
    done_count = 0
    for fut in as_completed(future_to_key):
        key = future_to_key[fut]
        res = fut.result()
        results_map[key] = res
        done_count += 1
        if done_count % 25 == 0 or done_count == len(upload_tasks):
            print(f"  Uploaded {done_count}/{len(upload_tasks)} assets...")

# Save cache
json.dump(UPLOAD_CACHE, open(CACHE_PATH, "w", encoding="utf-8"), indent=2)

# Attach Cloudinary URLs back to all_figures and all_frs
fig_url_map = {}
for fig in all_figures:
    up = results_map[fig["id"]]
    fig["imageUrl"] = up["url"]
    fig["imageWidth"] = up["width"]
    fig["imageHeight"] = up["height"]
    fig["imageBytes"] = up["bytes"]
    fig_url_map[fig["id"]] = up["url"]

    if "subImagesLocal" in fig and fig["subImagesLocal"]:
        fig["subImages"] = [
            {
                "id": sub_id,
                "label": sub_label,
                "imageUrl": results_map[sub_id]["url"],
                "width": results_map[sub_id]["width"],
                "height": results_map[sub_id]["height"],
            }
            for sub_id, sub_label, _ in fig["subImagesLocal"]
        ]
        del fig["subImagesLocal"]
    if "localPath" in fig:
        del fig["localPath"]

for fr in all_frs:
    fr["screenImageUrl"] = fig_url_map.get(fr["figureId"], "")
    fr["sequenceImageUrl"] = fig_url_map.get(fr["sequenceFigureId"], "")
    fr["classDiagramImageUrl"] = fig_url_map.get(fr["classDiagramFigureId"], "")

# Build final portalData.json
portal_data = {
    "projectInfo": {
        "name": "BrandHub — AI-Powered Multi-Channel Media Campaign Platform",
        "course": "SEP490 — Capstone Project (FPT University)",
        "cloudinaryCloudName": CLOUD_NAME,
        "totalFrs": len(all_frs),
        "totalUseCases": len(use_cases_table),
        "totalBusinessRules": len(business_rules),
        "totalMessages": len(messages_list),
        "totalFigures": len(all_figures),
        "canonicalRoles": CANONICAL_ROLES,
        "supportedPlatforms": ["Facebook", "Instagram", "TikTok", "Threads"],
        "paymentGateway": "SePay / PayOS (VietQR)",
    },
    "reports": [
        {
            "id": "report1",
            "code": "Report 1",
            "title": "Project Introduction",
            "fileName": "Report1_Project Introduction.docx",
            "description": "Background, competitor analysis (Hootsuite, Buffer, Jira, Notion), business opportunity, software product vision, and project scope across 9 domains.",
            "figureCount": len([f for f in all_figures if f["reportId"] == "report1"]),
        },
        {
            "id": "report2",
            "code": "Report 2",
            "title": "Project Management Plan",
            "fileName": "Report2_Project Management Plan.docx",
            "description": "Work Breakdown Structure (WBS), Agile-Scrum process with parallel AI iterations, 5-stage quality management, responsibility matrix, and polyrepo Git strategy.",
            "figureCount": len([f for f in all_figures if f["reportId"] == "report2"]),
        },
        {
            "id": "report3",
            "code": "Report 3",
            "title": "Software Requirement Specification (SRS)",
            "fileName": "report3_official.docx",
            "description": "Complete specification of 6 canonical actors, 141 atomic CRUD Use Cases (UC-001 – UC-141), 150 Official Figures (Context, Use Case, 6 Role Screen Flows, ERD, and 141 Screen Layouts), 65 Business Rules, and 60 Application Messages.",
            "figureCount": len([f for f in all_figures if f["reportId"] == "report3"]),
        },
        {
            "id": "report4",
            "code": "Report 4",
            "title": "Software Design Document (SDD — UML Models)",
            "fileName": "Report4_Software Design Document.docx / BrandHub-Master-All-In-One.vpp",
            "description": "9 Domain UML Class Diagrams (3.2 – 3.10) and 141 UML Sequence Diagrams mapped 1-to-1 with the 141 Functional Requirements.",
            "figureCount": len([f for f in all_figures if f["reportId"] == "report4"]),
        },
    ],
    "domains": DOMAINS,
    "actors": actors_table,
    "useCases": use_cases_table,
    "screenDescriptions": screen_descriptions,
    "screenAuthorization": screen_auth,
    "nonScreenFunctions": non_screen_funcs,
    "entities": entities_table,
    "businessRules": business_rules,
    "messages": messages_list,
    "frs": all_frs,
    "figures": all_figures,
}

out_json_path = os.path.join(DATA_DIR, "portalData.json")
with open(out_json_path, "w", encoding="utf-8") as f:
    json.dump(portal_data, f, ensure_ascii=False, indent=2)

print(f"Successfully wrote {out_json_path} ({os.path.getsize(out_json_path) / 1024:.1f} KB)!")
